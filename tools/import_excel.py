# ĐTCL Coach v1.0 (2026-10-06) — chuyển file Excel hướng dẫn (v1.8) của anh Peach thành data/*.json
"""
Chạy:  python tools/import_excel.py [đường_dẫn_file.xlsx]
Mặc định đọc tools/raw/excel/v1.8_2026-10-06_DTCL_Huong_Dan_Thach_Dau_Mua18.xlsx (giá trị đã tính sẵn trong file).
Ghi ra:
  data/hd_doi_hinh.json   19 đội hình (trang CHỌN ĐỘI HÌNH / XẾP HẠNG) + số liệu VN/Trung/Hàn
  data/hd_an.json         21 Ấn + 19 đội của trang CHỌN ẤN + ma trận điểm/diễn giải
  data/hd_loi.json        LÕI ĂN TOP 1, LÕI NÂNG CẤP, LÕI → CÁCH ĐÁNH (9 loại lõi + lõi tộc hệ)
  data/hd_nhip.json       lộ trình theo lối chơi × vòng (TRỢ LÝ VÁN ĐẤU), lời khuyên theo chuỗi / máu
  data/hd_kien_thuc.json  trang chủ, tinh linh & vòng đấu, kinh tế, xếp vị trí, lộ trình thách đấu,
                          từ điển, nguồn, máy chủ, tộc hệ, lên đồ, tướng
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "tools" / "raw" / "excel" / "v1.8_2026-10-06_DTCL_Huong_Dan_Thach_Dau_Mua18.xlsx"
DATA = ROOT / "data"
EXCEL_VERSION = "v1.8 · 06/10/2026"


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return v.strip()
    return v


def rows_of(ws, r1, r2, c1, c2):
    out = []
    for r in range(r1, r2 + 1):
        out.append([clean(ws.cell(r, c).value) for c in range(c1, c2 + 1)])
    return out


def col(letter: str) -> int:
    return openpyxl.utils.column_index_from_string(letter)


def split_list(s: str) -> list[str]:
    """Tách theo dấu phẩy / chấm giữa, bỏ qua dấu phẩy nằm trong ngoặc."""
    out, buf, depth = [], "", 0
    for ch in str(s or ""):
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        if ch in ",·" and depth == 0:
            out.append(buf)
            buf = ""
        else:
            buf += ch
    out.append(buf)
    return [x.strip() for x in out if x.strip()]


def header_map(ws, row=1):
    return {clean(ws.cell(row, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(row, c).value}


def table(ws, cols: list[str], first=2, last=None, key=None):
    h = header_map(ws)
    last = last or ws.max_row
    out = []
    for r in range(first, last + 1):
        row = {c: clean(ws.cell(r, h[c]).value) for c in cols}
        if key and not row.get(key):
            continue
        if not any(v not in ("", None) for v in row.values()):
            continue
        out.append(row)
    return out


def sheet_lines(ws):
    """Trả các dòng (danh sách ô có chữ) – dùng cho trang nội dung thuần chữ."""
    out = []
    for row in ws.iter_rows():
        vals = [clean(c.value) for c in row if c.value not in (None, "")]
        vals = [v for v in vals if not (isinstance(v, str) and v.startswith("◄"))]
        if vals:
            out.append(vals)
    return out


# --------------------------------------------------------------------------- đội hình
def build_doi_hinh(wb):
    ws = wb["DATA"]
    cols = ["name", "tier", "style", "diff", "cost", "rank", "traits", "overview", "level", "opener", "mid",
            "u1", "u2", "u3", "u4", "u5", "u6", "u7", "u8", "u9",
            "c1", "c1i1", "c1i2", "c1i3", "c1n", "c2", "c2i1", "c2i2", "c2i3", "c2n",
            "t", "ti1", "ti2", "ti3", "tn", "prio", "augs", "when", "pos", "tips", "weak", "tpl"]
    raw = table(ws, cols, 2, 40, key="name")
    out = []
    for r in raw:
        out.append({
            "ten": r["name"], "hang": r["tier"], "loi_choi": r["style"], "do_kho": r["diff"], "nhom_gia": r["cost"],
            "thu_tu": r["rank"], "toc_he": r["traits"], "tong_quan": r["overview"], "len_cap": r["level"],
            "mo_man": r["opener"], "giua_tran": r["mid"],
            "tuong": [r[f"u{i}"] for i in range(1, 10) if r[f"u{i}"]],
            "chu_luc": {"tuong": r["c1"], "do": [r["c1i1"], r["c1i2"], r["c1i3"]], "ghi_chu": r["c1n"]},
            "chu_luc_phu": {"tuong": r["c2"], "do": [r["c2i1"], r["c2i2"], r["c2i3"]], "ghi_chu": r["c2n"]},
            "do_don": {"tuong": r["t"], "do": [r["ti1"], r["ti2"], r["ti3"]], "ghi_chu": r["tn"]},
            "uu_tien_manh": r["prio"], "loi_khuyen_dung": split_list(r["augs"]), "khi_nao": r["when"],
            "xep_vi_tri": r["pos"], "meo": r["tips"], "diem_yeu": r["weak"], "mau_ban_co": r["tpl"],
        })
    out.sort(key=lambda x: x["thu_tu"] or 99)
    return out


def build_vi_tri(wb):
    ws = wb["DATA"]
    h = header_map(ws)
    res = []
    for r in range(2, 40):
        name = clean(ws.cell(r, h["pname"]).value)
        if not name:
            continue
        grid = [[clean(ws.cell(r, h[f"p{row * 7 + c}"]).value) for c in range(7)] for row in range(4)]
        res.append({"ten": name, "ghi_chu": clean(ws.cell(r, h["pnote"]).value), "luoi": grid})
    return res


def build_nhip(wb):
    ws = wb["DATA"]
    pace = table(ws, ["pkey", "plv", "pgold", "pact", "plvnum", "pevent"], 2, 80, key="pkey")
    styles = {}
    for p in pace:
        st, rd = p["pkey"].split("|")
        styles.setdefault(st, []).append({"vong": rd, "cap": p["plv"], "vang": p["pgold"], "viec": p["pact"],
                                          "cap_so": p["plvnum"], "su_kien": p["pevent"]})
    strat = table(ws, ["strat", "stratcomps"], 2, 10, key="strat")
    chuoi = table(ws, ["streak", "streaktxt"], 2, 10, key="streak")
    mau = table(ws, ["hp", "hptxt"], 2, 10, key="hp")
    stages = [r["stage"] for r in table(ws, ["stage"], 2, 30, key="stage")]
    return {
        "loi_choi": [{"ten": s["strat"], "doi_hinh": [x.strip("• ").strip() for x in str(s["stratcomps"]).split("\n") if x.strip()],
                      "lo_trinh": styles.get(s["strat"], [])} for s in strat],
        "vong": stages,
        "chuoi": [{"ten": c["streak"], "loi_khuyen": c["streaktxt"]} for c in chuoi],
        "mau": [{"ten": m["hp"], "loi_khuyen": m["hptxt"]} for m in mau],
    }


def build_items_extra(wb):
    ws = wb["DATA"]
    items = table(ws, ["iname", "ivn", "irole", "ieff", "iusers", "irecipe"], 2, 80, key="iname")
    comps = table(ws, ["complabel", "compstat", "compfor"], 2, 20, key="complabel")
    units = table(ws, ["uname", "ucost", "utraits", "umeta"], 2, 100, key="uname")
    traits = table(ws, ["tname", "ttype", "tbp", "teff", "tunits"], 2, 80, key="tname")
    return {
        "trang_bi": [{"ten": i["iname"], "cong_thuc": i["irecipe"] or i["ivn"], "danh_cho": i["irole"],
                      "tac_dung": i["ieff"], "tuong_nen_cam": split_list(i["iusers"])} for i in items],
        "manh_do": [{"ten": c["complabel"], "chi_so": c["compstat"],
                     "doi_hinh": [x.strip("• ").strip() for x in str(c["compfor"]).split("\n") if x.strip()]} for c in comps],
        "tuong": [{"ten": u["uname"], "gia": u["ucost"], "toc_he": split_list(u["utraits"]), "doi_hinh": split_list(u["umeta"])}
                  for u in units],
        "toc_he": [{"ten": t["tname"], "loai": t["ttype"], "moc": t["tbp"], "hieu_ung": t["teff"], "tuong": t["tunits"]}
                   for t in traits],
    }


# --------------------------------------------------------------------------- ấn
def build_an(wb):
    ws = wb["DATA_AN"]
    em = []
    for r in range(2, 23):
        a = [clean(ws.cell(r, c).value) for c in range(1, 8)]
        if a[0]:
            em.append({"ten": a[0], "toc_he": a[1], "cong_thuc": a[2], "hieu_ung": a[3], "deo_tot_nhat": a[4],
                       "hang_tb": a[5], "loai": a[6]})
    comps = []
    for r in range(2, 21):
        g = lambda L: clean(ws.cell(r, col(L)).value)  # noqa: E731
        if not g("M"):
            continue
        comps.append({
            "ten": g("M"), "diem_goc": g("N"),
            "chu_luc": {"tuong": g("O"), "kieu": g("P"), "do": [x.strip() for x in str(g("Q")).split("·")]},
            "chu_luc_phu": {"tuong": g("R"), "kieu": g("S"), "do": [x.strip() for x in str(g("T")).split("·")]},
            "do_don": {"tuong": g("U"), "do": [x.strip() for x in str(g("V")).split("·")]},
            "tuong": split_list(g("W")), "toc_he": g("X"), "so_lieu_vn": g("Y"), "ban_moi": g("Z"), "len_cap": g("AA"),
            "lech_do": [g("AC"), g("AD"), g("AE"), g("AF")],
        })
    n_c, n_e = len(comps), len(em)

    def mat(r0, numeric=True):
        m = []
        for i in range(n_c):
            row = []
            for j in range(n_e):
                v = clean(ws.cell(r0 + i, 2 + j).value)
                row.append(float(v or 0) if numeric else str(v))
            m.append(row)
        return m

    loai_do = [clean(ws.cell(r, col("K")).value) for r in range(2, 6)]
    return {
        "nguon": f"Excel hướng dẫn {EXCEL_VERSION} (trang CHỌN ẤN → ĐỘI HÌNH)",
        "cach_tinh": "Điểm = sức mạnh đội theo số liệu máy chủ VN (đã cân cỡ mẫu) + dự báo bản 18.4 + điểm từ Ấn "
                     "(mở được mốc tộc/hệ = điểm lớn; chỉ cộng chỉ số = điểm nhỏ; ấn vô dụng ≈ 0) + lệch theo loại đồ đang cầm. "
                     "Ấn thứ 2, 3 cùng loại dùng bảng riêng và chia đều.",
        "an": em, "loai_do": loai_do, "doi_hinh": comps,
        "diem": [mat(32), mat(54), mat(76)],
        "dien_giai": [mat(98, False), mat(120, False), mat(142, False)],
    }


# --------------------------------------------------------------------------- lõi
def build_loi(wb):
    ws = wb["DATA_AN"]
    loai = []
    for r in range(300, 309):
        g = lambda L: clean(ws.cell(r, col(L)).value)  # noqa: E731
        loai.append({"ten": g("A"), "nhan": g("B"), "vi_du": g("C"), "len_cap": g("D"), "doi_tuong": g("E"),
                     "chuoi": g("F"), "doi_hinh": [x.strip() for x in str(g("G")).split("·") if x.strip()],
                     "loi_sau": g("H"), "tranh": g("I"),
                     "theo_bac": {"Bạc": g("J"), "Vàng": g("K"), "Kim Cương": g("L")},
                     "theo_vong": {"2-1": g("M"), "3-2": g("N"), "4-2": g("O")}})
    mau = [{"ten": clean(ws.cell(r, 1).value), "loi_khuyen": clean(ws.cell(r, 2).value)} for r in range(328, 331)]

    w = wb["LÕI → CÁCH ĐÁNH"]
    toc = []
    for r in range(30, 42):
        g = lambda L: clean(w.cell(r, col(L)).value)  # noqa: E731
        if g("C"):
            toc.append({"ten": re.sub(r"^\d+\.\s*", "", g("C")), "toc_he": g("E"), "suc_manh": g("F"), "cho": g("G"),
                        "doi_hinh": g("I"), "cach_danh": g("K"), "ban_moi": g("M")})
    w = wb["LÕI ĂN TOP 1"]
    top1_rules = []
    for r in range(7, 14):
        a, b = clean(w.cell(r, 2).value), clean(w.cell(r, 5).value)
        if a:
            top1_rules.append({"tieu_de": a, "noi_dung": b})
    top1 = []
    for r in range(17, 36):
        g = lambda L: clean(w.cell(r, col(L)).value)  # noqa: E731
        if g("C"):
            top1.append({"ten": g("C"), "uu_tien": g("E"), "vong": g("F"), "doi_hinh": g("G"), "lam_gi": g("I"), "ky_vong": g("J")})
    w = wb["LÕI NÂNG CẤP"]
    rules = []
    for r in range(7, 12):
        a, b = clean(w.cell(r, 2).value), clean(w.cell(r, 4).value)
        if a:
            rules.append({"moc": a, "quy_tac": b})
    strong = []
    for r in range(15, 55):
        g = lambda L: clean(w.cell(r, col(L)).value)  # noqa: E731
        if g("C"):
            strong.append({"ten": g("C"), "loai": g("D"), "hang": g("E"), "hop_voi": g("F"), "ghi_chu": g("G")})
    return {"nguon": f"Excel hướng dẫn {EXCEL_VERSION}", "loai_loi": loai, "theo_mau": mau, "loi_toc_he": toc,
            "top1_nguyen_tac": top1_rules, "top1": top1, "quy_tac_moc": rules, "loi_manh": strong,
            "ghi_chu_top1": clean(wb["LÕI ĂN TOP 1"]["B37"].value),
            "quy_tac_chung": clean(wb["LÕI → CÁCH ĐÁNH"]["B55"].value)}


# --------------------------------------------------------------------------- kiến thức
def kv_rows(ws, r1, r2, ka="B", va="E"):
    out = []
    for r in range(r1, r2 + 1):
        k, v = clean(ws[f"{ka}{r}"].value), clean(ws[f"{va}{r}"].value)
        if k or v:
            out.append({"k": k, "v": v})
    return out


def build_kien_thuc(wb, extra):
    home = wb["TRANG CHỦ"]
    nguyen_tac = []
    for r in range(41, 46):
        for c in ("B", "I"):
            v = clean(home[f"{c}{r}"].value)
            if v:
                nguyen_tac.append(re.sub(r"^\d+\.\s*", "", v))
    ban_moi = kv_rows(home, 48, 56)

    tl = wb["TINH LINH & VÒNG ĐẤU"]
    tl_rules = [re.sub(r"^✦\s*", "", clean(tl[f"B{r}"].value)) for r in range(7, 16) if tl[f"B{r}"].value]
    wisps = []
    for r in range(19, 38):
        g = lambda L: clean(tl[f"{L}{r}"].value)  # noqa: E731
        if g("C"):
            wisps.append({"ten": g("C"), "loai": g("D"), "gia": g("E"), "hieu_ung": g("F"), "khi_nao": g("G")})
    vong = []
    for r in range(41, 63):
        g = lambda L: clean(tl[f"{L}{r}"].value)  # noqa: E731
        if g("B"):
            vong.append({"vong": g("B"), "su_kien": g("C"), "viec": g("E")})

    kt = wb["KINH TẾ & CẤP"]
    kinh_te = kv_rows(kt, 28, 34)
    lo_trinh_cap = kv_rows(kt, 37, 42)

    vt = wb["XẾP VỊ TRÍ"]
    vt_rules = [re.sub(r"^\d+\.\s*", "", clean(vt[f"B{r}"].value)) for r in range(18, 28) if vt[f"B{r}"].value]

    lt = wb["LỘ TRÌNH THÁCH ĐẤU"]
    habits = [{"nhom": clean(lt[f"B{r}"].value), "thoi_quen": clean(lt[f"C{r}"].value)} for r in range(10, 30) if lt[f"C{r}"].value]
    errors = [{"loi": clean(lt[f"C{r}"].value), "sua": clean(lt[f"D{r}"].value)} for r in range(32, 40) if lt[f"C{r}"].value]
    rank = [{"muc": clean(lt[f"B{r}"].value), "trong_tam": clean(lt[f"C{r}"].value), "chi_so": clean(lt[f"D{r}"].value)}
            for r in range(43, 48) if lt[f"B{r}"].value]

    td = wb["TỪ ĐIỂN THUẬT NGỮ"]
    terms = [{"ten": clean(td[f"B{r}"].value), "tieng_anh": clean(td[f"C{r}"].value), "nghia": clean(td[f"D{r}"].value)}
             for r in range(8, 34) if td[f"B{r}"].value]
    stats = [{"ten": clean(td[f"B{r}"].value), "nghia": clean(td[f"D{r}"].value)} for r in range(36, 45) if td[f"B{r}"].value]

    ng = wb["NGUỒN & PHIÊN BẢN"]
    phien_ban = kv_rows(ng, 7, 12, "B", "C")
    nhat_ky = kv_rows(ng, 15, 18, "B", "C")
    nguon = [{"ten": clean(ng[f"B{r}"].value), "url": clean(ng[f"C{r}"].value)} for r in range(20, 41) if ng[f"C{r}"].value]

    sv = wb["MÁY CHỦ VN - TRUNG - HÀN"]
    def tbl(r1, r2, cols):
        out = []
        for r in range(r1, r2 + 1):
            vals = [clean(sv[f"{c}{r}"].value) for c in cols]
            if vals[1]:
                out.append(vals)
        return out
    vn = tbl(8, 20, "BCDEFGH")
    cn = tbl(47, 51, "BCDEFG")
    kr = tbl(55, 68, "BCDEFG")
    so_sanh = []
    for r in range(30, 42, 2):
        so_sanh.append({"doi_hinh": clean(sv[f"B{r}"].value), "vn": clean(sv[f"D{r}"].value), "han": clean(sv[f"F{r}"].value),
                        "trung": clean(sv[f"H{r}"].value), "y_nghia": clean(sv[f"B{r + 1}"].value)})
    builds = []
    r = 75
    while r < sv.max_row:
        title = clean(sv[f"B{r}"].value)
        if title and not title.startswith(("🔑", "Giai", "✅", "⚔", "🛡", "💡")) and sv[f"F{r}"].value:
            b = {"ten": title, "so_lieu": clean(sv[f"F{r}"].value), "buoc": []}
            rr = r + 1
            while rr < sv.max_row and clean(sv[f"B{rr}"].value) and not (sv[f"F{rr}"].value and not str(sv[f"B{rr}"].value).startswith(("🔑", "Giai", "✅", "⚔", "🛡", "💡"))):
                k = clean(sv[f"B{rr}"].value)
                if k.startswith(("⚔", "🛡")):
                    v = " · ".join(str(clean(sv[f"{c}{rr}"].value)) for c in "EGI" if sv[f"{c}{rr}"].value)
                else:
                    v = clean(sv[f"D{rr}"].value)
                b["buoc"].append({"k": k, "v": v})
                rr += 1
            builds.append(b)
            r = rr
        else:
            r += 1
    may_chu = {"vn": {"tieu_de": clean(sv["B6"].value), "cot": ["#", "Đội hình", "Bậc", "Hạng TB", "Top 4", "Tỉ lệ chơi", "Tướng chính"],
                      "dong": vn, "goi_y": clean(sv["B22"].value)},
               "trung": {"cot": ["#", "Đội hình", "Hạng TB", "Top 4", "Tỉ lệ chơi", "Tướng chính"], "dong": cn},
               "han": {"cot": ["#", "Đội hình", "Hạng TB", "Top 4", "Tỉ lệ chơi", "Tướng chính"], "dong": kr},
               "so_sanh": so_sanh, "ghi_chu": clean(sv["B70"].value), "cach_build": builds}

    ld = wb["LÊN ĐỒ"]
    tao_tac = []
    for r in range(37, 49):
        g = lambda L: clean(ld[f"{L}{r}"].value)  # noqa: E731
        if g("C"):
            tao_tac.append({"ten": g("C"), "loai": g("E"), "tuong": g("G"), "ghi_chu": g("K")})
    do_rules = [re.sub(r"^\d+\.\s*", "", clean(ld[f"B{r}"].value)) for r in range(51, 57) if ld[f"B{r}"].value]

    return {
        "nguon_excel": f"Excel hướng dẫn {EXCEL_VERSION}",
        "nguyen_tac_vang": nguyen_tac, "ban_moi": ban_moi,
        "tinh_linh": {"quy_tac": tl_rules, "danh_sach": wisps}, "vong_dau": vong,
        "kinh_te": kinh_te, "lo_trinh_cap": lo_trinh_cap, "xep_vi_tri": vt_rules,
        "thoi_quen": habits, "loi_thuong_gap": errors, "ke_hoach_rank": rank,
        "tu_dien": terms, "chi_so": stats, "phien_ban": phien_ban, "nhat_ky_excel": nhat_ky, "nguon": nguon,
        "may_chu": may_chu, "tao_tac": tao_tac, "nguyen_tac_do": do_rules, **extra,
    }


def main(path: Path = SRC):
    wb = openpyxl.load_workbook(path, data_only=True)
    extra = build_items_extra(wb)
    files = {
        "hd_doi_hinh.json": {"nguon": f"Excel hướng dẫn {EXCEL_VERSION}", "doi_hinh": build_doi_hinh(wb), "vi_tri": build_vi_tri(wb)},
        "hd_an.json": build_an(wb),
        "hd_loi.json": build_loi(wb),
        "hd_nhip.json": build_nhip(wb),
        "hd_kien_thuc.json": build_kien_thuc(wb, extra),
    }
    for name, obj in files.items():
        (DATA / name).write_text(json.dumps(obj, ensure_ascii=False, indent=1), encoding="utf-8")
        print(name, (DATA / name).stat().st_size)


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else SRC)
